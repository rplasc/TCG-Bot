import asyncio
import aiosqlite
from openpyxl import load_workbook

async def insert_from_excel(file_path):
    wb = load_workbook(file_path, data_only=True)
    sheet = wb.active

    async with aiosqlite.connect("data/cards.db") as db:
        for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            name, rarity, attack, defense, hp, image_url, collection_id = row[:7]
            
            if not name:
                continue

            print("Adding: " + name)

            cursor = await db.execute("SELECT id FROM collections WHERE name = 'Default'")
            result = await cursor.fetchone()
            default_collection_id = result[0] if result else None

            await db.execute(
                "INSERT INTO cards (name, rarity, attack, defense, hp, image, collection_id) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (name, rarity, attack, defense, hp, image_url, collection_id or default_collection_id)
            )
        await db.commit()

asyncio.run(insert_from_excel("cards.xlsx"))
