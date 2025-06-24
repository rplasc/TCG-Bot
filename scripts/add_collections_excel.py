import asyncio
import aiosqlite
from openpyxl import load_workbook

async def insert_collections_from_excel(file_path):
    wb = load_workbook(file_path)
    sheet = wb.active

    async with aiosqlite.connect("data/cards.db") as db:
        for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):  # Skip header row
            name = row[0]
            if name:
                print("Adding: " + name)

                try:
                    await db.execute("INSERT INTO collections (name) VALUES (?)", (name,))
                except aiosqlite.IntegrityError:
                    print(f"Collection '{name}' already exists. Skipping.")
        await db.commit()

asyncio.run(insert_collections_from_excel("collections.xlsx"))
